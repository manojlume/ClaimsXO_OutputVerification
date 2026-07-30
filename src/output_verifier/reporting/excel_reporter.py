"""Excel verification report generation using openpyxl."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from output_verifier.constants import VerificationStatus, VerificationType
from output_verifier.models import ExecutionSummary, ReportConfig, ReportRow
from output_verifier.utilities.filename_utils import build_run_artifact_filename
from output_verifier.utilities.path_utils import ensure_directory

LOGGER = logging.getLogger(__name__)

DETAIL_HEADERS = [
    "Verification Type",
    "Scenario",
    "Run Type",
    "Output Type",
    "Expected File",
    "Actual File",
    "Status",
    "Issue",
]

COLUMN_WIDTHS = {
    1: 18,
    2: 36,
    3: 12,
    4: 14,
    5: 40,
    6: 40,
    7: 14,
    8: 80,
}


def build_report_filename(
    report_config: ReportConfig,
    verification_type: VerificationType,
    immediate_parent: str,
    timestamp: datetime,
) -> str:
    """
    Build the Excel report filename.

    Pattern:
      VerificationReport_<verification-type>_<immediate parent>_<yyyyMMddHHmmss>.xlsx
    """
    timestamp_token = timestamp.strftime(report_config.timestamp_format)
    return build_run_artifact_filename(
        prefix=report_config.base_file_name,
        verification_type=verification_type.value,
        immediate_parent=immediate_parent,
        timestamp=timestamp_token,
        extension=".xlsx",
    )


def _worksheet_name(verification_type: VerificationType) -> str:
    if verification_type == VerificationType.GDI:
        return "GDI Verification"
    if verification_type == VerificationType.HSAB:
        return "HSAB Verification"
    return f"{verification_type.value} Verification"


def _status_fill(status: VerificationStatus, colors: dict[str, str]) -> PatternFill:
    color_key = {
        VerificationStatus.PASSED: "passed",
        VerificationStatus.FAILED: "failed",
        VerificationStatus.NOT_VERIFIED: "not_verified",
    }[status]
    hex_color = colors.get(color_key, "#FFFFFF").lstrip("#")
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _write_detail_sheet(
    workbook: Workbook,
    verification_type: VerificationType,
    rows: list[ReportRow],
    report_config: ReportConfig,
) -> Worksheet:
    sheet = workbook.active
    sheet.title = _worksheet_name(verification_type)

    header_font = Font(bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    top_alignment = Alignment(vertical="top")

    for column_index, header in enumerate(DETAIL_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = header_font
        cell.alignment = top_alignment

    for row_index, report_row in enumerate(rows, start=2):
        values = [
            report_row.verification_type,
            report_row.scenario,
            report_row.run_type,
            report_row.output_type,
            report_row.expected_file,
            report_row.actual_file,
            report_row.status.value,
            report_row.issue,
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if column_index == 8:
                cell.alignment = wrap_alignment
            else:
                cell.alignment = top_alignment
            if column_index == 7:
                cell.fill = _status_fill(report_row.status, report_config.status_colors)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    return sheet


def _write_summary_sheet(workbook: Workbook, summary: ExecutionSummary) -> Worksheet:
    sheet = workbook.create_sheet("Summary")
    header_font = Font(bold=True)

    rows = [
        ("Verification type", summary.verification_type.value),
        (
            "Execution start time",
            summary.execution_start_time.strftime("%Y-%m-%d %H:%M:%S"),
        ),
        (
            "Execution end time",
            summary.execution_end_time.strftime("%Y-%m-%d %H:%M:%S")
            if summary.execution_end_time
            else "",
        ),
        ("Expected root", str(summary.expected_root)),
        ("Actual root", str(summary.actual_root)),
        ("Total report rows", summary.total_rows),
        ("Passed count", summary.passed_count),
        ("Failed count", summary.failed_count),
        ("NotVerified count", summary.not_verified_count),
        ("Report path", str(summary.report_path) if summary.report_path else ""),
        ("Exit code", summary.exit_code),
    ]

    for row_index, (label, value) in enumerate(rows, start=1):
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        label_cell.font = header_font
        sheet.cell(row=row_index, column=2, value=value)

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 80
    return sheet


class ExcelReporter:
    """Write verification results to an Excel workbook."""

    def write_report(
        self,
        report_config: ReportConfig,
        summary: ExecutionSummary,
        *,
        immediate_parent: str,
        timestamp: datetime,
    ) -> Path:
        """Create the Excel report and return its path."""
        ensure_directory(report_config.output_folder)
        filename = build_report_filename(
            report_config,
            summary.verification_type,
            immediate_parent,
            timestamp,
        )
        report_path = report_config.output_folder / filename

        LOGGER.info("Writing Excel report to %s", report_path)
        workbook = Workbook()
        _write_detail_sheet(
            workbook,
            summary.verification_type,
            summary.report_rows,
            report_config,
        )
        summary.report_path = report_path
        _write_summary_sheet(workbook, summary)

        try:
            workbook.save(report_path)
        except OSError as error:
            LOGGER.error("Unable to save Excel report %s: %s", report_path, error)
            raise

        LOGGER.info("Excel report created: %s", report_path)
        return report_path
