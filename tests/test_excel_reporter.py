"""Tests for Excel report generation."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from output_verifier.constants import SolutionType, VerificationStatus, VerificationType
from output_verifier.models import ExecutionSummary, ReportConfig, ReportRow
from output_verifier.reporting.excel_reporter import ExcelReporter, build_report_filename


def test_build_report_filename() -> None:
    config = ReportConfig(
        output_folder=Path("."),
        base_file_name="VerificationReport",
        include_solution_type=True,
        include_tpa_id=True,
        include_timestamp=True,
        timestamp_format="%Y%m%d%H%M%S",
        status_colors={},
    )
    name = build_report_filename(
        config,
        SolutionType.CLAIMS_XO,
        VerificationType.GDI,
        "20260612_000022_fullrun",
        datetime(2026, 7, 30, 11, 30, 45),
    )
    assert (
        name
        == "VerificationReport_ClaimsXO_GDI_20260612_000022_fullrun_20260730113045.xlsx"
    )


def test_excel_reporter_creates_sheets(tmp_path: Path) -> None:
    report_config = ReportConfig(
        output_folder=tmp_path,
        base_file_name="VerificationReport",
        include_solution_type=True,
        include_tpa_id=True,
        include_timestamp=True,
        timestamp_format="%Y%m%d%H%M%S",
        status_colors={
            "passed": "#C6EFCE",
            "failed": "#FFC7CE",
            "not_verified": "#FFEB9C",
        },
    )
    summary = ExecutionSummary(
        solution_type=SolutionType.CLAIMS_XO,
        verification_type=VerificationType.GDI,
        execution_start_time=datetime(2026, 7, 30, 11, 0, 0),
        execution_end_time=datetime(2026, 7, 30, 11, 30, 45),
        expected_root=tmp_path / "expected",
        actual_root=tmp_path / "actual" / "20260612_000022_fullrun",
        report_rows=[
            ReportRow(
                solution_type="ClaimsXO",
                verification_type="GDI",
                scenario="s1",
                run_type="current",
                output_type="LH1",
                expected_file="GDI_LH1_20260730110000.xml",
                actual_file="GDI_LH1_20260730120000.xml",
                status=VerificationStatus.PASSED,
            )
        ],
        exit_code=0,
    )

    path = ExcelReporter().write_report(
        report_config,
        summary,
        immediate_parent="20260612_000022_fullrun",
        timestamp=datetime(2026, 7, 30, 11, 30, 45),
    )
    assert (
        path.name
        == "VerificationReport_ClaimsXO_GDI_20260612_000022_fullrun_20260730113045.xlsx"
    )
    workbook = load_workbook(path)
    assert "ClaimsXO GDI Verification" in workbook.sheetnames
    assert "Summary" in workbook.sheetnames
    detail = workbook["ClaimsXO GDI Verification"]
    assert detail.cell(1, 1).value == "Solution Type"
    assert detail.cell(2, 8).value == "Passed"
